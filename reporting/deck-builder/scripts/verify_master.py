# -*- coding: utf-8 -*-
"""原生母版几何校验：画布/比例、无整页图充版面、形状越界、文本估行溢出、每行汉字数。

用法:
  python verify_master.py [pptx]          # 独立命令行
  deck.py check --pptx X --mode native    # 经 deck.py 调用 verify()，同一套判据
"""
import os
import sys

# Windows 控制台默认 GBK，中文输出在管道/重定向里会乱码（不影响判据，只影响可读性）
if sys.stdout.encoding and sys.stdout.encoding.lower() not in ("utf-8", "utf8"):
    sys.stdout.reconfigure(encoding="utf-8")

from pptx import Presentation
from pptx.util import Emu
from PIL import ImageFont

FONT_PATH = r"C:\Windows\Fonts\msyh.ttc"
EMU_PT = 12700.0
CM_PT = 72 / 2.54


def cm(emu):
    return emu / 360000.0


def width_pt(text, size):
    f = ImageFont.truetype(FONT_PATH, int(size))
    return f.getlength(text)


def wrap_lines(text, size, box_w_pt):
    """按字符贪心折行，返回行数。"""
    if not text:
        return 1
    lines, cur = 1, 0.0
    for ch in text:
        w = width_pt(ch, size)
        if cur + w > box_w_pt and cur > 0:
            lines += 1
            cur = 0.0
        cur += w
    return lines


def verify(pptx_path: str) -> int:
    """原生母版校验，返回 0=OK / 1=FAILED。判据与独立命令行完全一致。"""
    prs = Presentation(pptx_path)
    W, H = prs.slide_width, prs.slide_height
    CONTENT_BOTTOM = 26.5 * CM_PT   # 正文区下沿
    FOOT_RULE = 27.52 * CM_PT       # 页脚分隔线
    rc = 0

    print(f"[verify] {os.path.basename(pptx_path)}  canvas {cm(W):.2f}x{cm(H):.2f} cm  "
          f"ratio {W/H:.4f}  {os.path.getsize(pptx_path)/1024:.1f} KB  {len(prs.slides)} slides")
    if abs(W / H - 16 / 9) / (16 / 9) > 0.005:
        print("  ! canvas ratio off 16:9")
        rc = 1

    # ---- 排版容量参考（规格第 8 节字号：正文18 表14 图注12 导语20 结论条14；只印数字，不进判据） ----
    for size, box_cm, label in ((18, 24.0, "栏宽 24.0cm 正文"), (18, 48.64, "版心 48.64cm 正文"),
                                (14, 48.64, "表格/结论条 14pt"), (12, 48.64, "图注 12pt"),
                                (20, 24.0, "导语 20pt")):
        cw = width_pt("测", size)
        print(f"[verify] {size}pt 汉字宽 {cw:.2f}pt @ {label}: 每行约 {int(box_cm*CM_PT//cw)} 字")

    for idx, slide in enumerate(prs.slides, 1):
        pics = [sh for sh in slide.shapes if sh.shape_type == 13]
        full = [p for p in pics
                if p.width >= W * 0.98 and p.height >= H * 0.98 and p.top < 36000 and p.left < 36000]
        # 本页是否有页脚分隔线（y=27.52cm 的细横条）
        has_footer = any(abs(sh.top - FOOT_RULE * EMU_PT) < 20000 and sh.height < 40000
                         for sh in slide.shapes)
        over, overflow = [], []
        for sh in slide.shapes:
            if sh.left < 0 or sh.top < 0 or sh.left + sh.width > W + 1000 or sh.top + sh.height > H + 1000:
                over.append(f"{sh.shape_type}:{sh.name} "
                            f"({cm(sh.left):.2f},{cm(sh.top):.2f},{cm(sh.width):.2f}x{cm(sh.height):.2f})")
            # 文本估行
            if sh.has_text_frame and sh.text_frame.text.strip():
                bh = 0.0
                for p in sh.text_frame.paragraphs:
                    if not p.text.strip():
                        continue
                    size = max([(r.font.size.pt if r.font.size else 16) for r in p.runs] or [16])
                    sp = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
                    n = wrap_lines(p.text, size, sh.width / EMU_PT)
                    bh += n * size * 1.2 * sp
                    if p.space_after:
                        bh += p.space_after.pt
                bottom_pt = sh.top / EMU_PT + bh
                if not has_footer:
                    limit = H / EMU_PT                      # 无页脚的页（封面/章首页/结尾页）
                elif sh.top / EMU_PT > FOOT_RULE:
                    limit = H / EMU_PT                      # 页脚行本身在分隔线下方
                else:
                    limit = CONTENT_BOTTOM
                if bottom_pt > limit + 2:
                    overflow.append(f"{sh.name} text-bottom {bottom_pt:.0f}pt > limit {limit:.0f}pt")
        flag = "OK"
        if full:
            flag = "FULL-BLEED-PIC"
            rc = 1
        if over:
            flag += " OUT-OF-CANVAS"
            rc = 1
        if overflow:
            flag += " TEXT-OVERFLOW"
            rc = 1
        print(f"  L{idx}: shapes={len(slide.shapes)} pictures={len(pics)} full_bleed={len(full)} {flag}")
        for m in over:
            print("     !", m)
        for m in overflow:
            print("     !", m)

    print("[verify]", "FAILED" if rc else "OK")
    return rc


if __name__ == "__main__":
    PPTX = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "report-master-6up.pptx")
    raise SystemExit(verify(PPTX))
