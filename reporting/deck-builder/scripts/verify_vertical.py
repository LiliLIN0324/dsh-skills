# -*- coding: utf-8 -*-
"""竖版 9:16 母版几何校验（1080x1920，规格 _docs/master-layout-spec.md 第 6/9 节）。

用法: python verify_vertical.py <pptx> [<pptx> ...]
检查项：画布比例 9:16、无整页图充版面、形状不出血、文本按字体度量估行后不越界/不压页脚、
规格第 9 节版式判据（一屏视觉主体 ≤1、并排 ≤2 张且宽≥420 或高≥840、图内联数 image_cols）。
"""
import os
import re
import sys
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from pptx import Presentation
from PIL import ImageFont

FONT_PATH = r"C:\Windows\Fonts\msyh.ttc"
PX = 9525.0          # EMU per px
FONT_CACHE = {}

TOP_H = 560
BODY_TOP, BODY_BOT = 584, 1144      # 主体内容框 936x560
POINTS_TOP = 1208
FOOTER_TOP = 1830
CANVAS = (1080, 1920)


def f(size):
    k = int(size)
    if k not in FONT_CACHE:
        FONT_CACHE[k] = ImageFont.truetype(FONT_PATH, k)
    return FONT_CACHE[k]


def wrap_lines(text, size_pt, box_w_px):
    """size_pt 是 pt，box 用 px 表示（1pt = 1/0.75 px）。"""
    if not text:
        return 1
    px_size = size_pt / 0.75
    box_w_pt = box_w_px * 0.75
    ff = f(px_size)
    lines, cur = 1, 0.0
    for ch in text:
        w = ff.getlength(ch)
        if cur + w > box_w_pt and cur > 0:
            lines += 1
            cur = 0.0
        cur += w
    return lines


def image_cols(pic):
    """读机器标注 image_cols=N（图内并排 N 联）；无标注返回 None。

    约定（10-19 设计顾问裁定）：写 cNvPr@title（shape._element.nvPicPr.cNvPr.set("title",
    "image_cols=N")），descr/alt 留给屏幕阅读器的人类描述，不许放机器元数据。
    过渡期兼容读 descr，以后只认 title。
    YAML front matter 的 image_cols 由构建脚本落到 title，校验器判 936/N ≥ 420。
    """
    try:
        cNvPr = pic._element.nvPicPr.cNvPr
    except Exception:
        return None
    m = re.search(r"image_cols\s*=\s*(\d+)", cNvPr.get("title") or cNvPr.get("descr") or "")
    return int(m.group(1)) if m else None


def check(path):
    prs = Presentation(path)
    W, H = prs.slide_width, prs.slide_height
    w_px, h_px = round(W / PX), round(H / PX)
    bad = []
    if (w_px, h_px) != CANVAS:
        bad.append(f"canvas {w_px}x{h_px} != {CANVAS[0]}x{CANVAS[1]}")
    if abs(w_px / h_px - 9 / 16) > 1e-4:
        bad.append(f"ratio {w_px/h_px:.4f} != 0.5625")

    print(f"[verify-v] {os.path.basename(path)}  {w_px}x{h_px}px ratio {w_px/h_px:.4f}  "
          f"{os.path.getsize(path)/1048576:.2f} MB  {len(prs.slides)} slides")
    for i, slide in enumerate(prs.slides, 1):
        # 内容页 = 有 0..560px 全宽品牌块；封面色块结构不同，只查画布越界
        is_content = any(sh.top == 0 and sh.width == W and round(sh.height / PX) == TOP_H
                         for sh in slide.shapes)
        pics = [sh for sh in slide.shapes if sh.shape_type == 13]
        full = [p for p in pics if p.width >= W * 0.9 and p.height >= H * 0.9]
        # 视觉主体聚类：媒体图（排除整页图与页首全宽品牌块）。
        # 同主体 = 纵向重叠并排，或同宽全宽竖叠（规格 9：同宽全宽竖叠算 1 个主体）。
        media = [p for p in pics if p not in full
                 and not (p.top == 0 and p.width >= W * 0.98)]
        FULLW = 0.85 * 936          # ≥795px 视作全宽图
        subs = []
        for p in sorted(media, key=lambda s: (s.top, s.left)):
            for sub in subs:
                q = sub[0]
                ov = min(p.top + p.height, q.top + q.height) - max(p.top, q.top)
                if ov > 0.5 * min(p.height, q.height) or (p.width / PX >= FULLW and q.width / PX >= FULLW):
                    sub.append(p)
                    break
            else:
                subs.append([p])
        layout = []
        if len(subs) > 1:
            layout.append(f"一屏 {len(subs)} 个视觉主体（规格第9节：≤1）")
        for sub in subs:
            if len(sub) > 2:
                layout.append(f"并排 {len(sub)} 张（规格第9节：≤2，第 3 张拆页）")
            elif len(sub) == 2:
                # 放宽（规格 9）：宽 ≥420 或 高 ≥840，满足其一即放行；手机竖截 388×840 合规
                narrow = [f"{p.name} {round(p.width / PX)}x{round(p.height / PX)}px"
                          for p in sub if p.width / PX < 420 and p.height / PX < 840]
                if narrow:
                    layout.append("并排图过小(需宽≥420或高≥840): " + ", ".join(narrow))
        # 图内联数（人标注 image_cols=N）：判 936/N ≥ 420，机器看不了图内容
        for p in media:
            n = image_cols(p)
            if n and 936 / n < 420:
                layout.append(f"图内 {n} 联过密（936/{n}={936 / n:.0f}px < 420）: {p.name}")
        over, overflow, clash = [], [], []
        for sh in slide.shapes:
            x, y = sh.left / PX, sh.top / PX
            r, b = x + sh.width / PX, y + sh.height / PX
            if x < -1 or y < -1 or r > w_px + 1 or b > h_px + 1:
                over.append(f"{sh.shape_type}:{sh.name} box=({x:.0f},{y:.0f},{r:.0f},{b:.0f})")
            if sh.has_text_frame and sh.text_frame.text.strip():
                txt_h = 0.0
                for p in sh.text_frame.paragraphs:
                    if not p.text.strip():
                        continue
                    size = max([(run.font.size.pt if run.font.size else 22) for run in p.runs] or [22])
                    sp = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
                    n = wrap_lines(p.text, size, sh.width / PX)
                    txt_h += n * (size / 0.75) * 1.2 * sp
                bottom = y + txt_h
                if bottom > h_px + 1:
                    overflow.append(f"{sh.name} text-bottom {bottom:.0f}px > {h_px}px")
                if not is_content:
                    continue
                # 主体框内文字不得压到要点区；要点区文字不得压页脚
                if y < BODY_BOT <= bottom and y >= BODY_TOP:
                    clash.append(f"{sh.name} text-bottom {bottom:.0f}px 越过主体区下沿 {BODY_BOT}px")
                if POINTS_TOP <= y < FOOTER_TOP <= bottom:
                    clash.append(f"{sh.name} text-bottom {bottom:.0f}px 压页脚 {FOOTER_TOP}px")
                if FOOTER_TOP > y >= BODY_TOP and bottom > FOOTER_TOP:
                    clash.append(f"{sh.name} text-bottom {bottom:.0f}px 越过页脚 {FOOTER_TOP}px")
        flag = "OK" if not (full or over or overflow or clash or layout) else "FAIL"
        print(f"  P{i:02d}: shapes={len(slide.shapes)} pics={len(pics)} full_bleed={len(full)} "
              f"subjects={len(subs)} {flag}")
        for m in full:
            print("     ! FULL-BLEED-PIC", m.name)
        for m in over + overflow + clash + layout:
            print("     !", m)
        bad += over + overflow + clash + layout
        if full:
            bad.append("full-bleed picture")
    return 1 if bad else 0


if __name__ == "__main__":
    rc = 0
    for p in sys.argv[1:]:
        rc |= check(p)
    print("[verify-v]", "FAILED" if rc else "OK")
    raise SystemExit(rc)
